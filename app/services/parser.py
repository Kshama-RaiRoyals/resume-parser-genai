"""Extract readable text from supported resume formats."""
from io import BytesIO
from pathlib import Path
import fitz
from docx import Document


SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
MAX_FILE_BYTES = 10 * 1024 * 1024


class ResumeParseError(ValueError):
    """Raised when a file cannot be parsed or is unsupported."""


def extract_text(filename: str, content: bytes) -> str:
    """Extract text from PDF or DOCX bytes, validating file type and content."""
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Upload a PDF or DOCX file.")
    if not content:
        raise ResumeParseError("The uploaded file is empty.")
    if len(content) > MAX_FILE_BYTES:
        raise ResumeParseError("File exceeds the 10 MB upload limit.")

    try:
        if suffix == ".pdf":
            with fitz.open(stream=content, filetype="pdf") as pdf:
                text = "\n".join(page.get_text("text") for page in pdf)
        else:
            document = Document(BytesIO(content))
            paragraphs = [p.text for p in document.paragraphs if p.text.strip()]
            table_text = [
                " | ".join(cell.text.strip() for cell in row.cells)
                for table in document.tables
                for row in table.rows
            ]
            text = "\n".join(paragraphs + table_text)
    except Exception as exc:
        raise ResumeParseError("Could not read this file. Please upload a valid PDF or DOCX.") from exc

    if not text.strip():
        raise ResumeParseError("No readable text found. Scanned/image-only PDFs are not supported.")
    return text.strip()
