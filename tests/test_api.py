"""API integration tests; use MOCK_LLM to avoid network calls."""
import os
os.environ["MOCK_LLM"] = "true"

from fastapi.testclient import TestClient
from app.main import app, RESUME_STORE
from docx import Document
from io import BytesIO


def make_docx() -> bytes:
    buffer = BytesIO()
    doc = Document()
    doc.add_paragraph("Aarav Kulkarni")
    doc.add_paragraph("aarav.kulkarni@example.com")
    doc.add_paragraph("Python FastAPI SQL Generative AI Streamlit")
    doc.save(buffer)
    return buffer.getvalue()


def test_health():
    with TestClient(app) as client:
        response = client.get("/health")
        assert response.status_code == 200
        assert response.json() == {"status": "ok"}


def test_upload_docx_and_retrieve():
    RESUME_STORE.clear()
    with TestClient(app) as client:
        response = client.post(
            "/api/upload",
            files={"file": ("aarav_cv.docx", make_docx(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )
        assert response.status_code == 201, response.text
        body = response.json()
        assert body["document_id"]
        assert body["extracted_data"]["contact_information"]["email"] == "aarav.kulkarni@example.com"
        assert body["mode"] == "mock"
        retrieved = client.get(f"/api/resume/{body['document_id']}")
        assert retrieved.status_code == 200
        assert retrieved.json()["document_id"] == body["document_id"]


def test_unsupported_extension():
    with TestClient(app) as client:
        response = client.post("/api/upload", files={"file": ("resume.txt", b"not a resume", "text/plain")})
        assert response.status_code == 415


def test_empty_pdf_rejected():
    with TestClient(app) as client:
        response = client.post("/api/upload", files={"file": ("empty.pdf", b"", "application/pdf")})
        assert response.status_code == 422


def test_missing_document_returns_404():
    with TestClient(app) as client:
        response = client.get("/api/resume/not-a-real-id")
        assert response.status_code == 404
